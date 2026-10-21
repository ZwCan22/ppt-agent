from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from app.schemas.outline import PptOutline


def render_pptx(outline: PptOutline, output_path: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    blue = RGBColor(24, 49, 83)
    white = RGBColor(255, 255, 255)
    gray = RGBColor(55, 65, 81)

    for draft in outline.slides:
        special = draft.kind in ("cover", "closing")
        centered = special and not draft.bullets

        # 5是只有标题的版式，1是标题和正文版式
        layout = presentation.slide_layouts[5 if centered else 1]
        slide = presentation.slides.add_slide(layout)

        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = blue if special else white

        title = slide.shapes.title
        title.left = Inches(0.8)
        title.top = Inches(2.1 if centered else 0.6)
        title.width = Inches(11.7)
        title.height = Inches(3.0 if centered else 1.2)
        title.text = draft.title
        title.text_frame.word_wrap = True
        title.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

        for paragraph in title.text_frame.paragraphs:
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(42 if centered else 32)
            paragraph.font.bold = True
            paragraph.font.color.rgb = white if special else blue
            paragraph.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT

        if draft.bullets:
            body = slide.placeholders[1]
            body.left = Inches(0.8)
            body.top = Inches(2.0)
            body.width = Inches(11.7)
            body.height = Inches(4.7)

            frame = body.text_frame
            frame.clear()
            frame.word_wrap = True
            frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

            for index, bullet in enumerate(draft.bullets):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = bullet
                paragraph.level = 0
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(24)
                paragraph.font.color.rgb = white if special else gray
                paragraph.space_after = Pt(14)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output_path))
    return output_path